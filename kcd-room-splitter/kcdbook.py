"""kcdbook.py - fill the engineering pricing workbook (the `mike_p_for_services` template family) from a price
run, WITHOUT re-saving the workbook.

The template .xltm (or a filled .xlsm) is copied part for part. Only these parts change:
  [Content_Types].xml      the main part becomes a macro-enabled workbook instead of a template
  xl/workbook.xml          calcPr gets fullCalcOnLoad="1" so Excel recomputes every sheet when the file opens
  xl/sharedStrings.xml     the new texts are appended (count/uniqueCount updated)
  the Data Entry sheet     the answer cells (B1 date, B5 client number, B7 name, B9 room, F1 designer, F5 material,
                           ...) and the paste block from row 2: K price, L unit #, M part, N first size, O, P second size
  the order tabs           only when asked (--cell / --cells): the typed line items of Purchase Order, Doormark A1
                           and B1, VJT Order, CF5 Order, CF Thermo, CF Stone, Job Cart Cover Sheet, Final Price Sheet
VBA, ActiveX controls, the two external links, comments, drawings, printer settings and every other sheet stay
byte for byte (verify() proves it). Never re-save this workbook with openpyxl: it drops the links and controls.

Paste columns, as the engineers paste KCD's UNIT PRICE REVIEW (checked on the WERNER workbook): a part with two
sizes puts the first in N and the second in P (panels and shelves: depth, length; doors: width, height; backs:
width, height); a part with one size puts it in P; kicks, fillers, leg levelers and mirror backs show one size (the
width) in P although the price scripts know two. Column J (a shifted paste) and Q (an overflow of P) are read by
the Modify Here formulas; R is the engineer's own column. Nothing is written to J, Q or R.

Order tabs (checked on the WERNER workbook): their headers (PO date, client, room, engineer, install date, the
Doormark design and colour codes) are formulas that read Data Entry and are never written; only the line items
are typed cells. Rules for a typed cell on any tab:
  - a cell that holds a formula in the template is refused (the form computes it);
  - a cell inside a merged block is written at the block's top-left cell only (Doormark's boxes are merged blocks:
    design EM31, colour EM41, door rows R/Y/AR 55, 61, 67 ..., slab fillers CS/CZ/DT 166 ...);
  - the cell keeps the template's style; numbers go in as numbers, text through the shared-string table (the
    template has no inline strings), dates as serials and only into a cell whose format is a date;
  - a tick on the forms is the letter X in a cell (Doormark DN35 Match Back is preset; CO167 All Edges Square,
    the drawer styles, hinge sides are typed X's; no checkbox controls are involved); yes/no booleans are only
    accepted on a cell that already is a yes/no cell (there are none on the order tabs);
  - the Holes sheet is refused: it is the one sheet with a macro that runs when a cell changes.
Dropdown cells (Doormark design and colour, CF5 colour / insert / joint, VJT insert and colour, CF Thermo design)
are plain cells with a validation list; the exact list text has to be written (lists on Data Conversions).
"""
import re
import os
import sys
import json
import zipfile
import datetime
import html
import subprocess
import tempfile
import shutil

TEMPLATE_CT = 'application/vnd.ms-excel.template.macroEnabled.main+xml'
WORKBOOK_CT = 'application/vnd.ms-excel.sheet.macroEnabled.main+xml'
ONE_SIZE = re.compile(r'^(kick|filler|leg leveler|mirror back)\b', re.I)   # parts whose review line shows one size
NEW_ROW_STYLES = {'K': '380', 'L': '169', 'M': '169', 'N': '377', 'O': '377', 'P': '377'}   # as Excel wrote WERNER rows 43+
COL_LETTERS = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
MACRO_SHEETS = {'Holes': 'the Holes sheet runs a macro whenever one of its cells changes, so it has to be filled in Excel'}
DATE_FORMAT_IDS = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | set(range(50, 59))   # Excel's built-in date/time formats


def col_index(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def split_ref(ref):
    m = re.match(r'^([A-Z]+)(\d+)$', ref)
    if not m:
        raise ValueError('%r is not a cell address like B5 or AR55' % ref)
    return m.group(1), int(m.group(2))


def in_range(ref, first, last):
    """True when the cell is inside the block first:last."""
    c, r = split_ref(ref); c1, r1 = split_ref(first); c2, r2 = split_ref(last)
    return col_index(c1) <= col_index(c) <= col_index(c2) and r1 <= r <= r2


def is_date_format(code):
    """True when a number format shows a date or time: d, m, y, h or s outside quotes, brackets and escapes."""
    bare = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', '', code)
    return bool(re.search(r'[dmyhs]', bare, re.I))


def excel_serial(d):
    if isinstance(d, datetime.datetime):
        d = d.date()
    return (d - datetime.date(1899, 12, 30)).days


def xml_text(s):
    s = str(s)
    esc = s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    if s != s.strip() or '  ' in s:
        return '<t xml:space="preserve">%s</t>' % esc
    return '<t>%s</t>' % esc


# ----------------------------------------------------------------------------- the package
class Package:
    def __init__(self, path):
        self.path = path
        z = zipfile.ZipFile(path)
        self.order = [i.filename for i in z.infolist()]
        self.compress = {i.filename: i.compress_type for i in z.infolist()}
        self.parts = {n: z.read(n) for n in self.order}
        z.close()

    def text(self, name):
        return self.parts[name].decode('utf-8')

    def set_text(self, name, text):
        self.parts[name] = text.encode('utf-8')

    def sheet_part(self, sheet_name):
        wb = self.text('xl/workbook.xml')
        m = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(html.escape(sheet_name, quote=True)), wb)
        if not m:
            raise KeyError('no sheet named %r' % sheet_name)
        rels = self.text('xl/_rels/workbook.xml.rels')
        for rel in re.findall(r'<Relationship [^>]*/>', rels):
            if 'Id="%s"' % m.group(1) in rel:
                target = re.search(r'Target="([^"]+)"', rel).group(1)
                return 'xl/' + target.lstrip('/').replace('xl/', '', 1) if not target.startswith('/xl/') else target.lstrip('/')
        raise KeyError('sheet %r has no relationship' % sheet_name)

    def sheet_names(self):
        return [html.unescape(n) for n in re.findall(r'<sheet [^>]*name="([^"]*)"', self.text('xl/workbook.xml'))]

    def date_styles(self):
        """The cell style numbers (cellXfs index) whose number format is a date or time."""
        st = self.text('xl/styles.xml')
        custom = {int(i): html.unescape(c) for i, c in re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"', st)}
        xfs = re.search(r'<cellXfs[^>]*>(.*?)</cellXfs>', st, re.S).group(1)
        out = set()
        for i, xf in enumerate(re.findall(r'<xf [^>]*>', xfs)):
            m = re.search(r'numFmtId="(\d+)"', xf)
            fid = int(m.group(1)) if m else 0
            if fid in DATE_FORMAT_IDS or (fid in custom and is_date_format(custom[fid])):
                out.add(str(i))
        return out

    def write(self, out_path):
        with zipfile.ZipFile(out_path, 'w') as z:
            for n in self.order:
                z.writestr(zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0)), self.parts[n], compress_type=self.compress.get(n, zipfile.ZIP_DEFLATED))


# ----------------------------------------------------------------------------- shared strings
class SharedStrings:
    def __init__(self, pkg):
        self.pkg = pkg
        self.xml = pkg.text('xl/sharedStrings.xml')
        head = re.match(r'(.*?<sst[^>]*>)', self.xml, re.S).group(1)
        self.head = head
        self.items = re.findall(r'<si>.*?</si>', self.xml, re.S)
        self.tail = self.xml[self.xml.rfind('</sst>'):]
        self.plain = {}
        for i, it in enumerate(self.items):
            m = re.fullmatch(r'<si><t(?: xml:space="preserve")?>([^<]*)</t></si>', it, re.S)
            if m:
                self.plain.setdefault(html.unescape(m.group(1)), i)
        self.added = 0
        self.refs_added = 0

    def index(self, text):
        self.refs_added += 1
        if text in self.plain:
            return self.plain[text]
        self.items.append('<si>%s</si>' % xml_text(text))
        self.plain[text] = len(self.items) - 1
        self.added += 1
        return len(self.items) - 1

    def save(self):
        count = int(re.search(r'count="(\d+)"', self.head).group(1)) + self.refs_added
        head = re.sub(r'count="\d+"', 'count="%d"' % count, self.head, count=1)
        head = re.sub(r'uniqueCount="\d+"', 'uniqueCount="%d"' % len(self.items), head, count=1)
        self.pkg.set_text('xl/sharedStrings.xml', head + ''.join(self.items) + self.tail)


# ----------------------------------------------------------------------------- the sheet
class Sheet:
    """The cells of one worksheet part, edited in place; everything outside <sheetData> is kept verbatim."""

    def __init__(self, pkg, part, name='', new_styles=None, new_row_attrs=''):
        self.pkg, self.part, self.name = pkg, part, name
        self.new_styles = new_styles or {}               # style per column for a cell the template does not have (the paste rows)
        self.new_row_attrs = new_row_attrs
        self.xml = pkg.text(part)
        m = re.search(r'<sheetData>(.*?)</sheetData>|<sheetData/>', self.xml, re.S)
        if not m:
            raise ValueError('no sheetData in %s' % part)
        self.pre, self.post = self.xml[:m.start()], self.xml[m.end():]
        body = m.group(1) or ''
        self.rows = {}                                   # row number -> {'attrs': str, 'cells': {col: (attrs, body)}}
        for rm in re.finditer(r'<row r="(\d+)"([^>]*?)(?:/>|>(.*?)</row>)', body, re.S):
            r = int(rm.group(1)); cells = {}
            for cm in re.finditer(r'<c r="([A-Z]+)\d+"([^>]*?)(?:/>|>(.*?)</c>)', rm.group(3) or '', re.S):
                cells[cm.group(1)] = (cm.group(2) or '', cm.group(3) or '')
            self.rows[r] = {'attrs': rm.group(2) or '', 'cells': cells}
        self.formulas_before = len(re.findall(r'<f[ >/]', body))
        self.merges = re.findall(r'<mergeCell ref="([A-Z]+\d+):([A-Z]+\d+)"', self.post)
        self.col_styles = []                             # (first column, last column, style) from <cols>
        for ctag in re.findall(r'<col [^>]*/>', self.pre):
            a = dict(re.findall(r'(\w+)="([^"]*)"', ctag))
            if 'style' in a:
                self.col_styles.append((int(a['min']), int(a['max']), a['style']))
        self.date_styles = pkg.date_styles()
        self.written = []

    def style_for(self, col, row):
        """The cell's own style. For a cell the template does not have: the column's paste style (Data Entry K..P),
        else the row's own style, else the nearest styled cell above in the same column, else the column's style."""
        cur = self.rows.get(row, {}).get('cells', {}).get(col)
        if cur:
            m = re.search(r'\bs="(\d+)"', cur[0])
            if m:
                return m.group(1)
        if col in self.new_styles:
            return self.new_styles[col]
        if cur is not None:
            return None
        attrs = self.rows.get(row, {}).get('attrs', '')
        m = re.search(r'\bs="(\d+)"', attrs)
        if m and 'customFormat="1"' in attrs:
            return m.group(1)
        for r in range(row - 1, max(0, row - 200), -1):
            c = self.rows.get(r, {}).get('cells', {}).get(col)
            if c:
                m = re.search(r'\bs="(\d+)"', c[0])
                if m:
                    return m.group(1)
        n = col_index(col)
        for lo, hi, s in self.col_styles:
            if lo <= n <= hi:
                return s
        return None

    def set(self, ref, value, sst):
        """Write one typed cell. Refuses a formula cell, a cell inside a merged block that is not its top-left cell,
        a yes/no value on a cell that is not a yes/no cell, and a date on a cell whose format is not a date."""
        ref = ref.replace('$', '').upper()
        col, row = split_ref(ref)
        where = '%s!%s' % (self.name, ref) if self.name else ref
        if value is None or value == '':
            return
        cur = self.rows.get(row, {}).get('cells', {}).get(col)
        if cur and '<f' in cur[1]:
            f = re.search(r'<f[^>]*>([^<]*)</f>', cur[1])
            shown = ' (=%s)' % html.unescape(f.group(1))[:70] if f and f.group(1) else ''
            raise ValueError('%s holds a formula%s; the form computes that cell, so it is not written' % (where, shown))
        for first, last in self.merges:
            if first != ref and first != last and in_range(ref, first, last):
                raise ValueError("%s is inside the merged block %s:%s; write %s, the block's top-left cell" % (where, first, last, first))
        style = self.style_for(col, row)
        attrs = ' s="%s"' % style if style else ''
        if isinstance(value, bool):
            if not (cur and 't="b"' in cur[0]):
                raise ValueError('%s is not a yes/no cell; a tick on the order forms is the letter X' % where)
            attrs += ' t="b"'; inner = '<v>%d</v>' % (1 if value else 0)
        elif isinstance(value, datetime.date):
            if style not in self.date_styles:
                raise ValueError('%s is not a date cell (its number format is not a date), so %s would show as the number %d; '
                                 'write it as text or pick a date cell' % (where, value, excel_serial(value)))
            inner = '<v>%d</v>' % excel_serial(value)
        elif isinstance(value, (int, float)):
            v = ('%d' % value) if float(value).is_integer() else repr(float(value))
            inner = '<v>%s</v>' % v
        else:
            attrs += ' t="s"'; inner = '<v>%d</v>' % sst.index(str(value))
        self.rows.setdefault(row, {'attrs': self.new_row_attrs, 'cells': {}})
        self.rows[row]['cells'][col] = (attrs, inner)
        self.written.append((ref, value))

    def save(self):
        out = []
        for r in sorted(self.rows):
            cells = self.rows[r]['cells']
            parts = []
            for col in sorted(cells, key=col_index):
                attrs, body = cells[col]
                if body:
                    parts.append('<c r="%s%d"%s>%s</c>' % (col, r, attrs, body))
                else:
                    parts.append('<c r="%s%d"%s/>' % (col, r, attrs))
            out.append('<row r="%d"%s>%s</row>' % (r, self.rows[r]['attrs'], ''.join(parts)) if parts else '<row r="%d"%s/>' % (r, self.rows[r]['attrs']))
        body = ''.join(out)
        maxrow = max(self.rows) if self.rows else 1
        pre = re.sub(r'<dimension ref="([A-Z]+)1:([A-Z]+)\d+"/>', lambda m: '<dimension ref="%s1:%s%d"/>' % (m.group(1), m.group(2), maxrow), self.pre, count=1)
        formulas_after = len(re.findall(r'<f[ >/]', body))
        if formulas_after != self.formulas_before:
            raise RuntimeError('formula count changed %d -> %d' % (self.formulas_before, formulas_after))
        self.pkg.set_text(self.part, pre + '<sheetData>' + body + '</sheetData>' + self.post)


# ----------------------------------------------------------------------------- price lines
def read_price_lines(path):
    """The TSV kcdprice.py prints: price, unit, qty, part, dim1, dim2, order, script. Stops at the FLAGS block."""
    lines, flags = [], []
    for ln in open(path, encoding='utf-8', errors='replace'):
        ln = ln.rstrip('\n')
        if ln.startswith('price\t') or not ln.strip():
            continue
        if ln.startswith('FLAGS') or ln.startswith('  '):
            flags.append(ln.strip()); continue
        f = ln.split('\t') + [''] * 8
        try:
            lines.append({'price': float(f[0]), 'unit': int(f[1]), 'qty': int(f[2] or 1), 'part': f[3], 'dim1': f[4], 'dim2': f[5], 'order': f[6], 'script': f[7]})
        except ValueError:
            flags.append('unreadable line: ' + ln)
    return lines, flags


def paste_rows(lines, drop=None):
    """Turn engine lines into paste rows (K price, L unit, M part, N, O, P); a dropped line is listed with its reason."""
    rows, dropped = [], []
    for l in lines:
        hit = [why for rx, why in (drop or []) if rx.search('unit %d %s' % (l['unit'], l['part']))]
        if hit:
            dropped.append((l, hit[0])); continue
        d1, d2 = l['dim1'].strip(), l['dim2'].strip()
        if d1 and d2 and not ONE_SIZE.search(l['part']):
            n, p = d1, d2
        elif d1 and d2:
            n, p = '', d1
        elif d1 or d2:
            n, p = '', (d1 or d2)
        else:
            n, p = '', ''
        price = l['price']
        price = int(price) if float(price).is_integer() else round(price, 2)
        for _ in range(max(1, l['qty'])):
            rows.append({'K': price, 'L': l['unit'], 'M': l['part'], 'N': n, 'O': '', 'P': p})
    return rows, dropped


# ----------------------------------------------------------------------------- fill and verify
def fill(template, out_path, entries, rows, sheet_name='Data Entry', first_row=2, cells=None, skip_bad=False):
    """Copy the template to out_path with the Data Entry answers and paste rows, plus `cells` = {tab: {cell: value}}
    typed on any other tab. A cell that cannot be written (formula, merged block, bad tab, macro sheet, date on a
    non-date cell) stops the fill before anything is written, with every problem listed; with skip_bad the rest
    is written and the problems come back under 'refused'."""
    if os.path.abspath(out_path) == os.path.abspath(template):
        raise ValueError('refusing to overwrite the template')
    pkg = Package(template)
    ct = pkg.text('[Content_Types].xml')
    if TEMPLATE_CT in ct:
        pkg.set_text('[Content_Types].xml', ct.replace(TEMPLATE_CT, WORKBOOK_CT, 1))
    wb = pkg.text('xl/workbook.xml')
    if 'fullCalcOnLoad' not in wb:
        wb = re.sub(r'<calcPr ', '<calcPr fullCalcOnLoad="1" ', wb, count=1)
        pkg.set_text('xl/workbook.xml', wb)
    sst = SharedStrings(pkg)
    sheets, refused = {}, []

    def sheet(name):
        if name not in sheets:
            if name in MACRO_SHEETS:
                raise ValueError(MACRO_SHEETS[name])
            try:
                part = pkg.sheet_part(name)
            except KeyError:
                raise ValueError('there is no tab named "%s"; the tabs are: %s' % (name, ', '.join(pkg.sheet_names())))
            is_main = name == sheet_name
            sheets[name] = Sheet(pkg, part, name, NEW_ROW_STYLES if is_main else None, ' spans="11:16"' if is_main else '')
        return sheets[name]

    def put(name, ref, val):
        try:
            sheet(name).set(ref, val, sst)
        except ValueError as e:
            refused.append(str(e))

    for ref, val in entries.items():
        put(sheet_name, ref, val)
    for i, row in enumerate(rows):
        r = first_row + i
        for col in 'KLMNOP':
            if row.get(col) not in (None, ''):
                put(sheet_name, '%s%d' % (col, r), row[col])
    for name, d in (cells or {}).items():
        for ref, val in d.items():
            put(name, ref, val)
    if refused and not skip_bad:
        raise ValueError('\n'.join(refused))
    for sh in sheets.values():
        sh.save()
    sst.save()
    pkg.write(out_path)
    main_sheet = sheets.get(sheet_name)
    return {'cells': main_sheet.written if main_sheet else [], 'rows': len(rows), 'strings_added': sst.added,
            'sheet_part': pkg.sheet_part(sheet_name), 'sheet_parts': [sh.part for sh in sheets.values()],
            'written': {name: sh.written for name, sh in sheets.items()}, 'refused': refused}


def verify(template, out_path, sheet_parts):
    """Every part byte-identical to the template except content types, workbook.xml, the string table and the edited
    sheets; those parse; the string table only grew; every edited sheet keeps its formula count; the workbook loads.
    Returns a list of findings (empty = clean)."""
    import xml.dom.minidom
    if isinstance(sheet_parts, str):
        sheet_parts = [sheet_parts]
    a, b = Package(template), Package(out_path)
    findings = []
    if a.order != b.order:
        findings.append('part order differs')
    allowed = {'[Content_Types].xml', 'xl/workbook.xml', 'xl/sharedStrings.xml'} | set(sheet_parts)
    for n in a.order:
        if a.parts[n] != b.parts.get(n) and n not in allowed:
            findings.append('part changed: ' + n)
    for n in allowed:
        try:
            xml.dom.minidom.parseString(b.parts[n])
        except Exception as e:
            findings.append('%s not well-formed: %s' % (n, e))
    old = re.findall(rb'<si>.*?</si>', a.parts['xl/sharedStrings.xml'], re.S)
    new = re.findall(rb'<si>.*?</si>', b.parts['xl/sharedStrings.xml'], re.S)
    if new[:len(old)] != old:
        findings.append('the string table lost or changed an existing text')
    for n in sheet_parts:
        fa, fb = len(re.findall(rb'<f[ >/]', a.parts[n])), len(re.findall(rb'<f[ >/]', b.parts[n]))
        if fa != fb:
            findings.append('%s: formula count changed %d -> %d' % (n, fa, fb))
    try:
        import openpyxl
        wbk = openpyxl.load_workbook(out_path, read_only=True, keep_vba=True)
        wbk.close()
    except Exception as e:
        findings.append('openpyxl could not open the result: %s' % e)
    return findings


def preview(out_path, sheet='Final Price Sheet', rows=80):
    """What Excel will compute, approximately: LibreOffice converts a COPY to .xlsx (formulas recalculated,
    the two external links unresolved, no VBA) and the sheet's non-empty cells are returned as text lines."""
    tmp = tempfile.mkdtemp(prefix='kcdbook_')
    try:
        src = os.path.join(tmp, 'copy.xlsm'); shutil.copy(out_path, src)
        r = subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', tmp, src], capture_output=True, text=True, timeout=180)
        conv = os.path.join(tmp, 'copy.xlsx')
        if not os.path.exists(conv):
            return ['(LibreOffice could not convert the workbook: %s)' % (r.stderr or r.stdout).strip()[:300]]
        import openpyxl
        wbk = openpyxl.load_workbook(conv, data_only=True, read_only=True)
        if sheet not in wbk.sheetnames:
            return ['(no sheet %r; sheets: %s)' % (sheet, wbk.sheetnames)]
        ws = wbk[sheet]; out = []
        for row in ws.iter_rows(min_row=1, max_row=rows):
            vals = [(c.coordinate, c.value) for c in row if c.value not in (None, '')]
            if vals:
                out.append('  '.join('%s=%s' % (k, v) for k, v in vals))
        wbk.close()
        return out
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def parse_entry_value(text):
    t = text.strip()
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', t):
        y, m, d = map(int, t.split('-')); return datetime.date(y, m, d)
    if re.fullmatch(r'-?\d+', t):
        return int(t)
    if re.fullmatch(r'-?\d+\.\d+', t):
        return float(t)
    return t


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description='Fill the engineering pricing workbook from a kcdprice run without re-saving it.')
    ap.add_argument('template'); ap.add_argument('out')
    ap.add_argument('--lines', help='kcdprice.py output (TSV)')
    ap.add_argument('--drop', action='append', default=[], metavar='REGEX=REASON', help='leave out price lines where "unit N part" matches REGEX (repeatable), e.g. --drop "^unit (17|22) Door=no door on these units"')
    ap.add_argument('--entry', action='append', default=[], metavar='CELL=VALUE', help='Data Entry answer, e.g. B5=1091400, B1=2026-10-24, F5="Designer White Grain"')
    ap.add_argument('--entries', help='JSON file {cell: value}')
    ap.add_argument('--cell', action='append', default=[], metavar='TAB!CELL=VALUE', help='a typed cell on any tab (repeatable), e.g. "Purchase Order!C19=6", "Doormark A1!EM31=SARASOTA SHAKER"; a tick is X')
    ap.add_argument('--cells', help='JSON file {tab: {cell: value}} of typed cells on any tab')
    ap.add_argument('--report', help='write the plain-English report here')
    ap.add_argument('--preview', action='store_true', help='show the Final Price Sheet as LibreOffice computes it')
    a = ap.parse_args(argv)
    entries = {}
    if a.entries:
        for k, v in json.load(open(a.entries)).items():
            entries[k] = parse_entry_value(v) if isinstance(v, str) else v
    for s in a.entry:
        k, v = s.split('=', 1); entries[k.strip()] = parse_entry_value(v)
    cells = {}
    if a.cells:
        for tab, d in json.load(open(a.cells)).items():
            for k, v in d.items():
                cells.setdefault(tab, {})[k.strip()] = parse_entry_value(v) if isinstance(v, str) else v
    for s in a.cell:
        left, _, v = s.partition('=')
        tab, _, ref = left.rpartition('!')
        if not tab.strip() or not ref.strip():
            print('--cell needs TAB!CELL=VALUE, for example "Purchase Order!C19=6"; got: %s' % s); return 2
        cells.setdefault(tab.strip().strip("'"), {})[ref.strip()] = parse_entry_value(v)
    rows, dropped, flags = [], [], []
    if a.lines:
        lines, flags = read_price_lines(a.lines)
        drop = []
        for d in a.drop:
            rx, _, why = d.partition('=')
            drop.append((re.compile(rx, re.I), why or 'dropped'))
        rows, dropped = paste_rows(lines, drop)
    try:
        rep = fill(a.template, a.out, entries, rows, cells=cells)
    except ValueError as e:
        text = 'Nothing was written. Fix these and run again:\n  ' + str(e).replace('\n', '\n  ')
        print(text)
        if a.report:
            open(a.report, 'w', encoding='utf-8').write(text + '\n')
        return 2
    findings = verify(a.template, a.out, rep['sheet_parts'])
    out = ['Workbook %s from %s' % (os.path.basename(a.out), os.path.basename(a.template)), '']
    out += ['Data Entry answers written:'] + ['  %s = %s' % (k, v) for k, v in entries.items() if v not in (None, '')]
    out += ['', 'Price lines pasted in K..P from row 2: %d (%d new texts in the string table)' % (rep['rows'], rep['strings_added'])]
    tot = sum(r['K'] for r in rows if isinstance(r['K'], (int, float)))
    out += ['  the pasted prices add up to %.2f' % tot]
    for tab, written in rep['written'].items():
        if tab != 'Data Entry' and written:
            out += ['', 'Typed on the %s tab (%d cells):' % (tab, len(written))] + ['  %s = %s' % (ref, val) for ref, val in written]
    if dropped:
        out += ['', 'Price lines left out (%d):' % len(dropped)] + ['  %s unit %d %s %s %s: %s' % (l['price'], l['unit'], l['part'], l['dim1'], l['dim2'], why) for l, why in dropped]
    if flags:
        out += ['', 'Engine flags:'] + ['  ' + f for f in flags]
    out += ['', 'Check of the file: ' + ('every other part is byte for byte the template, the edited XML parses, the workbook opens' if not findings else '; '.join(findings))]
    if a.preview:
        out += ['', 'Final Price Sheet as LibreOffice computes it (no macros, external links unresolved):'] + ['  ' + s for s in preview(a.out)]
    text = '\n'.join(out)
    print(text)
    if a.report:
        open(a.report, 'w', encoding='utf-8').write(text + '\n')
    return 1 if findings else 0


if __name__ == '__main__':
    sys.exit(main())
